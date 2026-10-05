from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import load_workbook

from automations.conferencia_contas_receber import (
    BillingRow,
    ReceivablesResult,
    TotalCheck,
    billing_comparison,
    grouped,
    identify_ledger,
    save_excel,
)


class IdentifyLedgerTest(TestCase):
    header = ("DescricaoConta", "Debito", "Movimento", "Historico")

    def test_identifies_billing_ledger_by_content_regardless_of_filename(self):
        rows = [("Notas a Faturar", 10, 10, "Teste")]

        kind = identify_ledger(Path("relatorio qualquer.xlsx"), self.header, rows)

        self.assertEqual(kind, "razao_faturar")

    def test_identifies_commission_ledger_by_content_with_accents(self):
        rows = [("Comissão de Cartão de Crédito", 10, 10, "Teste")]

        kind = identify_ledger(Path("relatorio qualquer.xlsx"), self.header, rows)

        self.assertEqual(kind, "razao_comissao")

    def test_uses_normalized_filename_as_fallback(self):
        rows = [(None, 10, 10, "Teste")]

        billing = identify_ledger(
            Path("RAZÃO_notas-A-FATURAR.xlsx"), self.header, rows
        )
        commission = identify_ledger(
            Path("razao COMISSÃO cartão.xlsx"), self.header, rows
        )

        self.assertEqual(billing, "razao_faturar")
        self.assertEqual(commission, "razao_comissao")


class ClientGroupingTest(TestCase):
    header = ("Cliente", "Saldo")

    def test_consolidates_cvc_accounts_before_comparison(self):
        accounting = grouped(
            self.header,
            [
                ("CVC BRASIL", -107471.82),
                ("CVC BRASIL OPERADORA E AGENCIA", 12353.70),
                ("CVC BRASIL OPERADORA E AGENCIA DE VIAGENS", 146142.28),
            ],
            "Cliente",
            "Saldo",
        )
        financial = grouped(
            self.header,
            [("CVC Brasil Operadora e Agência de Viagens", 51024.16)],
            "Cliente",
            "Saldo",
        )

        self.assertEqual(accounting["CVC"], financial["CVC"])

    def test_does_not_merge_brt_and_bwt(self):
        result = grouped(
            self.header,
            [("BRT Operadora", 10), ("BWT Operadora", 20)],
            "Cliente",
            "Saldo",
        )

        self.assertEqual(result["BRT"][1], 10)
        self.assertEqual(result["BWTOPERADORA"][1], 20)

    def test_consolidates_related_trade_names(self):
        result = grouped(
            self.header,
            [
                ("DECOLAR.COM", 103700.37),
                ("DECOLAR.COM LTDA", 6499.08),
                ("Despegar-PAM", -3874.56),
            ],
            "Cliente",
            "Saldo",
        )

        self.assertEqual(result["DECOLARDESPEGAR"][1], Decimal("106324.89"))


class BillingDetailsTest(TestCase):
    def test_compares_bordero_transaction_with_ledger_sheet_number(self):
        bordero = (
            ("NumeroDaTransacao", "Valor", "Status"),
            [
                (101, Decimal("60"), "Baixado"),
                (101, Decimal("40"), "Baixado"),
                (202, Decimal("20"), "Baixado"),
            ],
        )
        ledger = (
            ("NumeroPlanilha", "Debito"),
            [("101", Decimal("100")), (303, Decimal("30"))],
        )

        rows = billing_comparison(bordero, ledger)

        by_identification = {row.identification: row for row in rows}
        self.assertEqual(by_identification["101"].status, "Conciliado")
        self.assertEqual(by_identification["202"].difference, Decimal("20"))
        self.assertEqual(by_identification["303"].difference, Decimal("-30"))
        self.assertEqual(
            sum((row.source_value for row in rows), Decimal()), Decimal("120")
        )
        self.assertEqual(
            sum((row.accounting_value for row in rows), Decimal()), Decimal("130")
        )

    def test_excel_contains_billing_details_sheet(self):
        billing_rows = [
            BillingRow("101", Decimal("100"), Decimal("100")),
            BillingRow("202", Decimal("20"), Decimal()),
        ]
        result = ReceivablesResult(
            clients=[],
            client_accounting_total=Decimal(),
            client_financial_total=Decimal(),
            billing=TotalCheck("Notas a faturar", Decimal("120"), Decimal("100")),
            commissions=TotalCheck(
                "Comissões de cartão", Decimal(), Decimal()
            ),
            billing_rows=billing_rows,
            hotel="Cumbuco",
        )

        with TemporaryDirectory() as directory:
            output = Path(directory) / "resultado.xlsx"
            save_excel(result, output)
            workbook = load_workbook(output, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames, ["Resumo", "Clientes", "Notas a faturar"]
                )
                sheet = workbook["Notas a faturar"]
                self.assertEqual(sheet["A4"].value, "Identificação (Transação / Nº planilha)")
                self.assertEqual(sheet["A5"].value, "101")
                self.assertEqual(sheet["E5"].value, "Conciliado")
                self.assertEqual(sheet["A6"].value, "202")
                self.assertEqual(sheet["D6"].value, 20)
                self.assertEqual(sheet["E6"].value, "Divergente")
            finally:
                workbook.close()
